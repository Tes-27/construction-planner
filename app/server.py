"""Construction Planner — เว็บแอปรันในเครื่อง

python app/server.py [ชื่อโครงการ] [--lan]   → เปิด http://localhost:5070
--lan = ให้มือถือในวง Wi-Fi เดียวกันเข้าได้ (ไม่มีรหัสผ่าน — ใช้เฉพาะเครือข่ายที่ไว้ใจ)
"""
import datetime as dt, io, os, sys, threading
from flask import Flask, jsonify, request, send_file, send_from_directory

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import core  # noqa: E402
import billing  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROJECTS = os.path.join(ROOT, "projects")
STATIC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")
app = Flask(__name__, static_folder=None)
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0  # ไม่ให้เบราว์เซอร์ใช้ไฟล์เก่าหลังอัปเดต
lock = threading.Lock()
state = {"name": None, "p": None}


def list_projects():
    return sorted(d for d in os.listdir(PROJECTS) if os.path.exists(os.path.join(PROJECTS, d, "project.json")))


def open_project(name):
    state["p"] = core.Project(os.path.join(PROJECTS, name))
    state["name"] = name


def P():
    return state["p"]


def qdate():
    s = request.args.get("date") or (request.get_json(silent=True) or {}).get("date")
    return core.parse_date(s) if s else dt.date.today()


def err(msg, code=400):
    return jsonify({"error": msg}), code


@app.route("/")
def index():
    return send_from_directory(STATIC, "index.html")


@app.route("/static/<path:f>")
def static_files(f):
    return send_from_directory(STATIC, f)


@app.route("/plans/<path:f>")
def plan_files(f):
    return send_from_directory(os.path.join(P().folder, "plans"), f)


@app.route("/api/projects")
def api_projects():
    return jsonify({"projects": list_projects(), "current": state["name"]})


@app.route("/api/open", methods=["POST"])
def api_open():
    name = request.get_json()["name"]
    if name not in list_projects():
        return err("ไม่พบโครงการ")
    with lock:
        open_project(name)
    return jsonify({"ok": True})


@app.route("/api/state")
def api_state():
    p = P()
    with lock:
        p.reload()
        val = p.validation()
    locs = [{**l, "value": p.loc_value.get(l["loc_id"], 0.0)} for l in p.locations]
    tasks = [{"task_id": t["task_id"], "name": t["name"], "trade": t["trade"], "steps": t["step_list"],
              "auto_time": t["auto_time"], "value": p.task_value.get(t["task_id"], 0.0)} for t in p.tasks]
    return jsonify({"project": state["name"], "meta": p.meta, "locations": locs, "tasks": tasks,
                    "total_value": p.total_value, "k_contract": p.k_contract, "validation": val,
                    "contract_value": p.contract_value, "contract_original": p.contract_original,
                    "baselines": p.list_baselines(), "today": dt.date.today().isoformat(),
                    "progress_last": max((r["date"] for r in p.progress_rows), default=None)})


@app.route("/api/schedule")
def api_schedule():
    p, d = P(), qdate()
    with lock:
        rows, sched, end = p.activity_rows(d)
        ms = p.milestone_status(d, sched, p.baseline())
        fin, eot_days = p.contract_finish()
    return jsonify({"activities": rows, "finish": p.cal.date(end).isoformat(), "milestones": ms,
                    "baseline": p.meta.get("active_baseline"), "contract_finish": fin.isoformat(),
                    "contract_finish_original": p.meta["finish_date"], "eot_days": eot_days,
                    "start": p.cal.date(0).isoformat()})


@app.route("/api/activity/<act_id>", methods=["POST"])
def api_activity(act_id):
    p = P()
    with lock:
        try:
            p.update_activity(act_id, request.get_json())
        except (KeyError, ValueError) as e:
            p.reload()
            return err(str(e))
    return jsonify({"ok": True})


@app.route("/api/baseline", methods=["POST"])
def api_baseline():
    name = (request.get_json() or {}).get("name") or ("baseline-" + dt.date.today().isoformat())
    name = "".join(c for c in name if c.isalnum() or c in "-_ .ก-๙")[:60]
    with lock:
        P().save_baseline(name, status=qdate())
    return jsonify({"ok": True, "name": name})


@app.route("/api/grid")
def api_grid():
    p, d = P(), qdate()
    floor = request.args.get("floor", "1")
    with lock:
        locs = [l for l in p.locations if l["floor"] == floor and p.loc_value.get(l["loc_id"])]
        tids = []
        for t in p.tasks:
            if any((t["task_id"], l["loc_id"]) in p.cells for l in locs):
                tids.append(t["task_id"])
        cells = {}
        for (t, l), v in p.cells.items():
            if t in tids and p.loc[l]["floor"] == floor:
                st = p.cell_state(t, l, d)
                if p.task[t]["auto_time"]:
                    st["pct"] = p.auto_time_pct(d)
                cells[f"{t}|{l}"] = {"value": v, **st}
    return jsonify({"floor": floor, "locations": locs, "tasks": tids, "cells": cells})


@app.route("/api/progress", methods=["POST"])
def api_progress():
    body = request.get_json()
    d = core.parse_date(body.get("date") or dt.date.today().isoformat())
    if d > dt.date.today():
        return err("บันทึกความคืบหน้าล่วงหน้าไม่ได้ (วันที่อยู่ในอนาคต)")
    with lock:
        try:
            P().append_progress(d.isoformat(), body["entries"], body.get("user", ""))
        except ValueError as e:
            return err(str(e))
    return jsonify({"ok": True})


@app.route("/api/history")
def api_history():
    p = P()
    t, l = request.args["task"], request.args["loc"]
    h = p.progress_hist.get((t, l), [])
    return jsonify([{"date": r[0], "pct": r[2], "steps": r[3], "note": r[4]} for r in h])


@app.route("/api/scurve")
def api_scurve():
    p, d = P(), qdate()
    trade = request.args.get("trade") or None
    with lock:
        return jsonify(p.scurve(d, trade))


@app.route("/api/billing/<no>", methods=["POST"])
def api_billing(no):
    """action: submit / approve / pay / reset · date · invoice_no · note · evidence_date"""
    body = request.get_json() or {}
    act, d = body.get("action"), body.get("date") or dt.date.today().isoformat()
    if core.parse_date(d) > dt.date.today():
        return err("วันที่อยู่ในอนาคต")
    p = P()
    with lock:
        cur = p.billing.get(no, {})
        if act == "submit":
            f = {"submitted": d, "invoice_no": body.get("invoice_no", ""), "note": body.get("note", ""),
                 "evidence_date": body.get("evidence_date") or d}
        elif act == "approve":
            if not cur.get("submitted"):
                return err("ต้องบันทึกการยื่นเบิกก่อน")
            f = {"approved": d}
        elif act == "pay":
            if not cur.get("submitted"):
                return err("ต้องบันทึกการยื่นเบิกก่อน")
            f = {"paid": d, "approved": cur.get("approved") or d}
        elif act == "reset":
            f = {k: "" for k in ("submitted", "approved", "paid", "invoice_no", "evidence_date", "note")}
        else:
            return err("action ไม่ถูกต้อง")
        try:
            p.save_billing(no, f)
        except KeyError:
            return err("ไม่พบงวด " + no)
    return jsonify({"ok": True})


@app.route("/api/billing_report/<no>")
def api_billing_report(no):
    d = request.args.get("date")
    with lock:
        try:
            r = billing.report(P(), no, core.parse_date(d) if d else None)
        except KeyError:
            return err("ไม่พบงวด " + no, 404)
    return jsonify(r)


@app.route("/api/billing_summary")
def api_billing_summary():
    with lock:
        return jsonify(billing.summary(P(), qdate()))


@app.route("/api/export/billing/<no>.xlsx")
def export_billing(no):
    d = request.args.get("date")
    with lock:
        r = billing.report(P(), no, core.parse_date(d) if d else None)
    return send_file(billing.to_xlsx(r), as_attachment=True,
                     download_name=f"{state['name']}-ขอเบิกงวด{no}-{r['evidence_date']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/billing/<no>")
def billing_page(no):
    return send_from_directory(STATIC, "billing.html")


@app.route("/api/plans")
def api_plans():
    return jsonify(P().plans)


@app.route("/api/heat")
def api_heat():
    p, d = P(), qdate()
    with lock:
        return jsonify(p.heat(request.args.get("floor", "1"), d, request.args.get("trade") or None,
                              request.args.get("task") or None))


@app.route("/api/locinfo")
def api_locinfo():
    p, d = P(), qdate()
    with lock:
        try:
            return jsonify(p.loc_detail(request.args["loc"], d))
        except KeyError:
            return err("ไม่พบพื้นที่", 404)


@app.route("/api/zones", methods=["POST"])
def api_zones():
    b = request.get_json() or {}
    with lock:
        try:
            P().save_zone(b["sheet"], b["slot"], b.get("polys"))
        except KeyError:
            return err("ไม่พบแบบแผ่นนี้")
    return jsonify({"ok": True})


# ─────────────── VO / EOT / ค่าปรับ
@app.route("/api/changes")
def api_changes():
    p, d = P(), qdate()
    with lock:
        return jsonify({"vo": p.vo, "eot": p.eot, "delay": p.delay(d), "contract_original": p.contract_original,
                        "vo_approved": p.vo_approved, "contract_value": p.contract_value,
                        "ld_per_day": p.meta.get("ld_per_day"), "ld_cap_pct": p.meta.get("ld_cap_pct", 0.10)})


@app.route("/api/vo", methods=["POST"])
def api_vo():
    b = request.get_json() or {}
    p = P()
    with lock:
        try:
            if b.get("delete"):
                p.delete_vo(b["no"])
            else:
                p.save_vo(b)
            p.reload()
            bad = [i for i in p.issues if i[0] == "error" and "VO" in i[1]]
        except (ValueError, KeyError) as e:
            p.reload()
            return err(str(e))
    return jsonify({"ok": True, "warnings": [i[1] for i in bad]})


@app.route("/api/eot", methods=["POST"])
def api_eot():
    b = request.get_json() or {}
    p = P()
    with lock:
        try:
            p.delete_eot(b["no"]) if b.get("delete") else p.save_eot(b)
        except (ValueError, KeyError) as e:
            return err(str(e))
        p.reload()
    return jsonify({"ok": True})


@app.route("/api/settings", methods=["POST"])
def api_settings():
    b = request.get_json() or {}
    p = P()
    allowed = ("ld_per_day", "ld_cap_pct", "contract_no", "owner_address", "contractor_address", "contractor_tax_id")
    with lock:
        for k in allowed:
            if k in b:
                v = b[k]
                if k in ("ld_per_day", "ld_cap_pct"):
                    v = float(v) if v not in ("", None) else None
                if v in ("", None):
                    p.meta.pop(k, None)
                else:
                    p.meta[k] = v
        p.save_meta()
    return jsonify({"ok": True})


# ─────────────── รูปถ่าย
@app.route("/api/photos")
def api_photos():
    p = P()
    return jsonify(p.photos_for(loc=request.args.get("loc") or None, task=request.args.get("task") or None))


@app.route("/api/photos", methods=["POST"])
def api_photo_upload():
    import uuid
    from PIL import Image, ImageOps
    p = P()
    loc, task = request.form.get("loc_id", ""), request.form.get("task_id", "")
    d = request.form.get("date") or dt.date.today().isoformat()
    if loc not in p.loc:
        return err("ไม่พบพื้นที่")
    files = request.files.getlist("files")
    if not files:
        return err("ไม่มีไฟล์")
    out_dir, th_dir = os.path.join(p.folder, "photos"), os.path.join(p.folder, "photos", "thumbs")
    os.makedirs(th_dir, exist_ok=True)
    saved = []
    for f in files:
        try:
            im = ImageOps.exif_transpose(Image.open(f.stream)).convert("RGB")
        except Exception:
            return err(f"ไฟล์ {f.filename} ไม่ใช่รูปภาพ")
        pid = dt.datetime.now().strftime("%Y%m%d-%H%M%S-") + uuid.uuid4().hex[:6]
        big = im.copy()
        big.thumbnail((1600, 1600))
        big.save(os.path.join(out_dir, pid + ".jpg"), "JPEG", quality=82)
        im.thumbnail((360, 360))
        im.save(os.path.join(th_dir, pid + ".jpg"), "JPEG", quality=75)
        rec = {"id": pid, "file": pid + ".jpg", "loc_id": loc, "task_id": task, "date": d,
               "note": request.form.get("note", ""), "uploaded_at": dt.datetime.now().isoformat(timespec="seconds"),
               "w": big.width, "h": big.height, "name": f.filename}
        with lock:
            p.add_photo(rec)
        saved.append(rec)
    return jsonify({"ok": True, "photos": saved})


@app.route("/api/photos/<pid>", methods=["DELETE"])
def api_photo_delete(pid):
    p = P()
    with lock:
        rec = p.delete_photo(pid)
    if rec:
        for path in (os.path.join(p.folder, "photos", rec["file"]), os.path.join(p.folder, "photos", "thumbs", rec["file"])):
            if os.path.exists(path):
                os.remove(path)
    return jsonify({"ok": bool(rec)})


@app.route("/photos/<path:f>")
def photo_files(f):
    return send_from_directory(os.path.join(P().folder, "photos"), f)


@app.route("/api/boq")
def api_boq():
    return jsonify(P().boq)


@app.route("/api/export/schedule.xlsx")
def export_schedule():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    p, d = P(), qdate()
    with lock:
        rows, sched, end = p.activity_rows(d)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gantt"
    head = ["รหัส", "กิจกรรม", "หมวด", "ชั้น", "ระยะเวลา (วันทำงาน)", "เริ่ม", "เสร็จ", "Baseline เริ่ม",
            "Baseline เสร็จ", "Float", "วิกฤต", "น้ำหนัก %", "มูลค่า BOQ", "% คืบหน้า", "Predecessors"]
    ws.append([p.meta["name"]])
    ws.append([f"ข้อมูล ณ {d.isoformat()} · แผนเสร็จ {p.cal.date(end).isoformat()} · สัญญาเสร็จ {p.meta['finish_date']}"])
    ws.append(head)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="334155")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([r["act_id"], r["name"], r["trade"], r["floors"], r["duration"], core.parse_date(r["start"]),
                   core.parse_date(r["finish"]), core.parse_date(r["b_start"]) if r["b_start"] else None,
                   core.parse_date(r["b_finish"]) if r["b_finish"] else None, r["float"],
                   "ใช่" if r["critical"] else "", round(r["weight"], 3), round(r["value"], 2), round(r["pct"], 1),
                   r["preds"]])
    for row in ws.iter_rows(min_row=4):
        for c in row[5:9]:
            c.number_format = "dd/mm/yyyy"
        row[12].number_format = "#,##0.00"
    widths = [8, 46, 8, 8, 10, 11, 11, 11, 11, 7, 6, 9, 14, 9, 30]
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    ws.freeze_panes = "C4"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{state['name']}-gantt-{d.isoformat()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    lan = "--lan" in sys.argv or os.environ.get("HOST") == "0.0.0.0"
    names = list_projects()
    want = args[0] if args else (names[0] if names else None)
    if not want or want not in names:
        sys.exit(f"ไม่พบโครงการ '{want}' ในโฟลเดอร์ projects/ (มี: {', '.join(names)})")
    open_project(want)
    port = int(os.environ.get("PORT", 5070))
    print(f"Construction Planner · โครงการ {want} · http://localhost:{port}")
    if lan:
        import socket
        try:
            sk = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            sk.connect(("8.8.8.8", 80))
            print(f"มือถือในวง Wi-Fi เดียวกัน: http://{sk.getsockname()[0]}:{port}")
            sk.close()
        except OSError:
            pass
    app.run(host="0.0.0.0" if lan else "127.0.0.1", port=port, debug=False, threaded=True)
