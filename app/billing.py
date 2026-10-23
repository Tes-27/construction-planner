"""ใบขอเบิกเงินงวด — สร้างข้อมูลรายงาน + Excel (ไม่มีโค้ดเฉพาะโครงการ)"""
import datetime as dt
import io

TH_DIGIT = ["", "หนึ่ง", "สอง", "สาม", "สี่", "ห้า", "หก", "เจ็ด", "แปด", "เก้า"]
TH_POS = ["", "สิบ", "ร้อย", "พัน", "หมื่น", "แสน"]
TH_MONTH = ["", "มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน", "กรกฎาคม", "สิงหาคม",
            "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]


def _read_int(n):
    """อ่านจำนวนเต็มเป็นภาษาไทย (รองรับหลักล้านซ้อน)"""
    if n == 0:
        return ""
    if n >= 1_000_000:
        return _read_int(n // 1_000_000) + "ล้าน" + _read_int(n % 1_000_000)
    out, digits = "", str(n)
    L = len(digits)
    for i, ch in enumerate(digits):
        d, pos = int(ch), L - i - 1
        if d == 0:
            continue
        if pos == 1 and d == 1:
            out += "สิบ"
        elif pos == 1 and d == 2:
            out += "ยี่สิบ"
        elif pos == 0 and d == 1 and L > 1:
            out += "เอ็ด"
        else:
            out += TH_DIGIT[d] + TH_POS[pos]
    return out


def bahttext(x):
    x = round(float(x) + 1e-9, 2)
    baht, satang = int(x), int(round((x - int(x)) * 100))
    if baht == 0 and satang == 0:
        return "ศูนย์บาทถ้วน"
    s = (_read_int(baht) + "บาท") if baht else ""
    return s + (_read_int(satang) + "สตางค์" if satang else "ถ้วน")


def thai_date(s):
    if not s:
        return ""
    d = dt.date.fromisoformat(s) if isinstance(s, str) else s
    return f"{d.day} {TH_MONTH[d.month]} {d.year + 543}"


FLOOR_ORDER = ["SITE", "FDN", "1", "2", "3", "4", "5", "6", "7", "RF"]
FLOOR_NAME = {"SITE": "ภายนอก", "FDN": "ฐานราก", "RF": "หลังคา"}
TRADE_NAME = {"PRELIM": "งานเตรียมการ", "STR": "งานโครงสร้าง", "ARCH": "งานสถาปัตยกรรม", "EE": "งานระบบไฟฟ้า",
              "SN": "งานระบบสุขาภิบาล", "FP": "งานดับเพลิง"}


def report(p, no, date=None):
    """ข้อมูลทั้งหมดของใบขอเบิกงวด `no` · date = วันที่หลักฐานผลงาน (ค่าเริ่ม: evidence_date ที่บันทึกไว้ หรือวันนี้)"""
    sched, _ = p.schedule(status=date or dt.date.today())
    ms_all = p.milestone_status(date or dt.date.today(), sched, p.baseline())
    m = next((x for x in ms_all if x["no"] == no), None)
    if m is None:
        raise KeyError(no)
    ev_date = date or (dt.date.fromisoformat(m["evidence_date"]) if m.get("evidence_date") else dt.date.today())
    if date is None and m.get("evidence_date"):  # ประเมินใหม่ ณ วันหลักฐาน
        sched, _ = p.schedule(status=ev_date)
        ms_all = p.milestone_status(ev_date, sched, p.baseline())
        m = next(x for x in ms_all if x["no"] == no)
    cv = p.contract_value  # รวม VO ที่อนุมัติแล้ว
    k = p.k_contract
    for c in m["checks"]:
        if c["task"] == "ALL":
            c["label"] = "งานทั้งโครงการ"
        else:
            fl = c["floors"]
            where = ("ชั้น " + fl) if fl[:1].isdigit() else FLOOR_NAME.get(fl, fl)
            c["label"] = f"{p.task[c['task']]['name']} · {where}" if c["task"] in p.task else c["text"]

    # ตารางสรุปการเบิกทุกงวด + ยอดสะสม
    rows, cum = [], {"amount": 0.0, "retention": 0.0, "wht": 0.0, "net": 0.0, "vat": 0.0}
    prev = {"amount": 0.0, "retention": 0.0, "wht": 0.0, "net": 0.0, "vat": 0.0}
    idx = [x["no"] for x in ms_all].index(no)
    for i, x in enumerate(ms_all):
        billed = x["status"] in ("submitted", "approved", "paid")
        this = x["no"] == no
        rows.append({"no": x["no"], "pct": x["pct"], "amount": x["amount"], "description": x["description"],
                     "status": "this" if this else x["status"], "submitted": x["submitted"], "paid": x["paid"],
                     "money": x["money"]})
        if (billed and i != idx) or this:
            for kk in cum:
                cum[kk] += x["money"][kk]
            if not this:
                for kk in prev:
                    prev[kk] += x["money"][kk]

    # ผลงานรายหมวดงาน (มูลค่าตามสัญญา) + แผน ณ วันหลักฐาน (Baseline ถ้ามี)
    base = p.baseline()
    plan_sched = base["activities"] if base else p.schedule()[0]
    n_idx = p.cal.idx_on_or_before(ev_date)
    plan_val = {}
    for a in p.activities:
        s = plan_sched.get(a["act_id"])
        if not s or not a["value"]:
            continue
        frac = max(0.0, min(1.0, (n_idx - s["es"]) / max(s["ef"] - s["es"] + 1, 1)))
        plan_val[a["task_id"]] = plan_val.get(a["task_id"], 0.0) + a["value"] * frac
    tasks = []
    for t in p.tasks:
        v = p.task_value.get(t["task_id"], 0.0)
        if not v:
            continue
        cells = {c: val for c, val in p.cells.items() if c[0] == t["task_id"]}
        earned = p.earned(ev_date, cells)
        tasks.append({"task_id": t["task_id"], "name": t["name"], "trade": t["trade"], "weight": 100 * v / p.total_value,
                      "value": v * k, "plan_pct": 100 * plan_val.get(t["task_id"], 0) / v,
                      "pct": 100 * earned / v, "earned": earned * k})
    trades = []
    for tr in ["PRELIM", "STR", "ARCH", "EE", "SN", "FP"]:
        ts = [t for t in tasks if t["trade"] == tr]
        if not ts:
            continue
        val = sum(t["value"] for t in ts)
        trades.append({"trade": tr, "name": TRADE_NAME.get(tr, tr), "tasks": ts, "value": val,
                       "weight": sum(t["weight"] for t in ts), "earned": sum(t["earned"] for t in ts),
                       "pct": 100 * sum(t["earned"] for t in ts) / val if val else 0,
                       "plan_pct": sum(t["plan_pct"] * t["value"] for t in ts) / val if val else 0})
    total_earned = sum(t["earned"] for t in tasks)
    total_plan = sum(t["plan_pct"] * t["value"] for t in tasks) / cv if cv else 0

    # ผลงานรายชั้น (หมวดงาน × ชั้น)
    floors = [f for f in FLOOR_ORDER if any(p.loc[l]["floor"] == f for (_, l) in p.cells)]
    matrix = []
    for t in p.tasks:
        if t["auto_time"] or not p.task_value.get(t["task_id"]):
            continue
        row = {"name": t["name"], "trade": t["trade"], "cells": {}}
        for f in floors:
            cells = {c: val for c, val in p.cells.items() if c[0] == t["task_id"] and p.loc[c[1]]["floor"] == f}
            v = sum(cells.values())
            if v:
                row["cells"][f] = 100 * p.earned(ev_date, cells) / v
        matrix.append(row)

    # รูปถ่ายประกอบ: ช่วงตั้งแต่หลักฐานงวดก่อนหน้า ถึงวันหลักฐานงวดนี้ · เฉพาะหมวด/ชั้นในเงื่อนไข (ถ้ามี)
    prev_ev = max((x.get("evidence_date") or "" for x in ms_all
                   if x["no"] != no and x["status"] in ("submitted", "approved", "paid")
                   and (x.get("evidence_date") or "") < ev_date.isoformat()), default="")
    want = [(c["task"], set(c["floors"].split(",")) if c["floors"] else None) for c in m["checks"] if c["task"] != "ALL"]
    pics = []
    for ph in p.photos_for(start=prev_ev or None, end=ev_date.isoformat()):
        fl = p.loc.get(ph["loc_id"], {}).get("floor")
        ok = not want or any((not ph.get("task_id") or ph["task_id"] == t) and
                             (fs is None or fl in _fs(fs)) for t, fs in want)
        if ok:
            loc = p.loc.get(ph["loc_id"], {})
            pics.append({**ph, "loc_name": loc.get("name", ""), "floor": fl,
                         "task_name": p.task[ph["task_id"]]["name"] if ph.get("task_id") in p.task else ""})
    mm = m["money"]
    return {
        "photos": pics[:12], "photos_total": len(pics), "photo_from": prev_ev,
        "meta": {k2: p.meta.get(k2) for k2 in ("name", "owner", "contractor", "location", "contract_value_ex_vat",
                                               "vat_rate", "retention_rate", "wht_rate", "start_date", "finish_date",
                                               "owner_address", "contractor_address", "contractor_tax_id",
                                               "contract_no")},
        "milestone": {k2: m[k2] for k2 in ("no", "pct", "amount", "description", "status", "submitted", "approved",
                                           "paid", "invoice_no", "note", "ready", "forecast")},
        "checks": m["checks"], "money": mm, "net_text": bahttext(mm["net"]), "gross_text": bahttext(mm["gross"]),
        "evidence_date": ev_date.isoformat(), "evidence_date_th": thai_date(ev_date),
        "doc_date": m["submitted"] or dt.date.today().isoformat(),
        "doc_date_th": thai_date(m["submitted"] or dt.date.today()),
        "start_th": thai_date(p.meta["start_date"]), "finish_th": thai_date(p.meta["finish_date"]),
        "rows": rows, "cum": cum, "prev": prev, "contract_value": cv,
        "contract_gross": cv * (1 + float(p.meta.get("vat_rate", 0.07))),
        "trades": trades, "earned_total": total_earned, "earned_pct": 100 * total_earned / cv if cv else 0,
        "plan_pct": total_plan, "floors": [{"id": f, "name": FLOOR_NAME.get(f, "ชั้น " + f)} for f in floors],
        "matrix": matrix, "draft": not m["ready"],
        "baseline_name": base["name"] if base else None,
    }


def _fs(parts):
    from core import floor_set
    out = set()
    for x in parts:
        out |= floor_set(x)
    return out


def summary(p, date):
    """สรุปการเงินทั้งโครงการ"""
    sched, _ = p.schedule(status=date)
    ms = p.milestone_status(date, sched, p.baseline())
    agg = {k: {"amount": 0.0, "vat": 0.0, "retention": 0.0, "wht": 0.0, "net": 0.0} for k in
           ("submitted", "approved", "paid", "ready", "planned")}
    for m in ms:
        a = agg[m["status"]]
        for k in a:
            a[k] += m["money"][k]
    billed = {k: agg["submitted"][k] + agg["approved"][k] + agg["paid"][k] for k in agg["paid"]}
    return {"by_status": agg, "billed": billed, "paid": agg["paid"],
            "outstanding": {k: agg["submitted"][k] + agg["approved"][k] for k in agg["paid"]},
            "ready_not_billed": agg["ready"], "contract": p.money(p.contract_value)}


def to_xlsx(r):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = openpyxl.Workbook()
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="E5E7EB")
    bold = Font(bold=True)
    m, mm, meta = r["milestone"], r["money"], r["meta"]

    ws = wb.active
    ws.title = "ใบขอเบิก"
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 20
    lines = [
        ("หนังสือขอเบิกเงินค่างวดงาน" + (" (ร่าง — เงื่อนไขยังไม่ครบ)" if r["draft"] else ""), None),
        (f"เลขที่ {m['invoice_no'] or '..........'}", f"วันที่ {r['doc_date_th']}"),
        (f"เรียน {meta['owner']}", None),
        (f"โครงการ {meta['name']}", None),
        (f"สถานที่ {meta.get('location') or ''}", None),
        (f"มูลค่าสัญญา {meta['contract_value_ex_vat']:,.2f} บาท (ไม่รวม VAT) · ระยะเวลา {r['start_th']} – {r['finish_th']}", None),
        ("", None),
        (f"ผู้รับจ้างได้ดำเนินงานงวดที่ {m['no']} ({m['pct']*100:.0f}% ของมูลค่าสัญญา) แล้วเสร็จตามเงื่อนไข ณ วันที่ {r['evidence_date_th']}", None),
        (f"รายละเอียดงวด: {m['description']}", None),
        ("จึงขอเบิกเงินค่างวดงาน ดังนี้", None),
    ]
    for a, b in lines:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    money_rows = [("(1) มูลค่างวดงาน (ไม่รวม VAT)", mm["amount"]), ("(2) ภาษีมูลค่าเพิ่ม 7%", mm["vat"]),
                  ("(3) รวมทั้งสิ้น", mm["gross"]), ("(4) หัก เงินประกันผลงาน 5% ของ (1)", -mm["retention"]),
                  ("(5) ยอดหลังหักประกันผลงาน", mm["after_retention"]),
                  ("(6) หัก ภาษี ณ ที่จ่าย 3% ของ (1)", -mm["wht"]), ("(7) ยอดชำระสุทธิ", mm["net"])]
    for a, b in money_rows:
        ws.append([a, b])
        c = ws.cell(ws.max_row, 2)
        c.number_format = "#,##0.00;(#,##0.00)"
        for cc in ws[ws.max_row][:2]:
            cc.border = box
    ws.cell(ws.max_row, 1).font = bold
    ws.cell(ws.max_row, 2).font = bold
    ws.append([f"({r['net_text']})"])
    ws.append([])
    ws.append(["สรุปการเบิกสะสม", None])
    ws.cell(ws.max_row, 1).font = bold
    for a, b in [("เบิกงวดก่อนหน้า (ไม่รวม VAT)", r["prev"]["amount"]), ("งวดนี้", mm["amount"]),
                 ("เบิกสะสมรวมงวดนี้", r["cum"]["amount"]),
                 ("คงเหลือตามสัญญา", r["contract_value"] - r["cum"]["amount"]),
                 ("เงินประกันผลงานสะสม (ผู้ว่าจ้างหักไว้)", r["cum"]["retention"])]:
        ws.append([a, b])
        ws.cell(ws.max_row, 2).number_format = "#,##0.00"
    ws.append([])
    ws.append(["ลงชื่อ ........................ ผู้รับจ้าง", "ลงชื่อ ........................ ผู้ควบคุมงาน"])
    ws.append([f"({meta['contractor']})", ""])
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    ws2 = wb.create_sheet("ตารางงวด")
    hdr = ["งวด", "%", "รายการ", "(1) มูลค่า", "(2) VAT", "(3) รวม", "(4) ประกัน 5%", "(5) หลังหักประกัน",
           "(6) ณ ที่จ่าย 3%", "(7) สุทธิ", "สถานะ", "ยื่น", "ได้รับเงิน"]
    ws2.append(hdr)
    stat = {"this": "งวดนี้", "paid": "ได้รับเงินแล้ว", "approved": "อนุมัติแล้ว", "submitted": "ยื่นแล้ว",
            "ready": "พร้อมเบิก", "planned": "ยังไม่ถึง"}
    for x in r["rows"]:
        q = x["money"]
        ws2.append([x["no"], x["pct"], x["description"], q["amount"], q["vat"], q["gross"], -q["retention"],
                    q["after_retention"], -q["wht"], q["net"], stat.get(x["status"], x["status"]), x["submitted"],
                    x["paid"]])
    for c in ws2[1]:
        c.font, c.fill, c.border = bold, head, box
    for row in ws2.iter_rows(min_row=2):
        row[1].number_format = "0%"
        for c in row[3:10]:
            c.number_format = "#,##0.00;(#,##0.00)"
    for col, w in zip("ABCDEFGHIJKLM", [7, 6, 60, 15, 13, 15, 14, 15, 13, 15, 13, 11, 11]):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet("เงื่อนไขงวด")
    ws3.append([f"ผลตรวจเงื่อนไขงวดที่ {m['no']} ณ วันที่ {r['evidence_date_th']}"])
    ws3.append(["เงื่อนไข", "ต้องการ %", "ผลงาน %", "ผ่าน"])
    for c in ws3[2]:
        c.font, c.fill = bold, head
    for c in r["checks"]:
        ws3.append([c["label"], c["need"], round(c["pct"], 1), "✓" if c["ok"] else "✗"])
    ws3.column_dimensions["A"].width = 60

    ws4 = wb.create_sheet("ผลงานรายหมวด")
    ws4.append([f"สรุปผลงานสะสม ณ วันที่ {r['evidence_date_th']}" + (f" · แผนตาม Baseline: {r['baseline_name']}" if r["baseline_name"] else "")])
    ws4.append(["หมวดงาน", "น้ำหนัก %", "มูลค่า (บาท)", "แผน %", "ผลงาน %", "มูลค่าผลงาน (บาท)"])
    for c in ws4[2]:
        c.font, c.fill = bold, head
    for tr in r["trades"]:
        ws4.append([tr["name"], tr["weight"] / 100, tr["value"], tr["plan_pct"] / 100, tr["pct"] / 100, tr["earned"]])
        for c in ws4[ws4.max_row]:
            c.font = bold
        for t in tr["tasks"]:
            ws4.append(["   " + t["name"], t["weight"] / 100, t["value"], t["plan_pct"] / 100, t["pct"] / 100, t["earned"]])
    ws4.append(["รวมทั้งโครงการ", 1, r["contract_value"], r["plan_pct"] / 100, r["earned_pct"] / 100, r["earned_total"]])
    for c in ws4[ws4.max_row]:
        c.font = bold
    for row in ws4.iter_rows(min_row=3):
        row[1].number_format = row[3].number_format = row[4].number_format = "0.0%"
        row[2].number_format = row[5].number_format = "#,##0"
    ws4.column_dimensions["A"].width = 44
    for col in "BCDEF":
        ws4.column_dimensions[col].width = 15

    ws5 = wb.create_sheet("ผลงานรายชั้น")
    ws5.append(["หมวดงาน"] + [f["name"] for f in r["floors"]])
    for c in ws5[1]:
        c.font, c.fill = bold, head
    for row in r["matrix"]:
        ws5.append([row["name"]] + [(row["cells"][f["id"]] / 100 if f["id"] in row["cells"] else None) for f in r["floors"]])
    for row in ws5.iter_rows(min_row=2):
        for c in row[1:]:
            c.number_format = "0%"
    ws5.column_dimensions["A"].width = 40
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
